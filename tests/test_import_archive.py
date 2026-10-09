"""Offline tests for import_archive.py. Workbooks and vault notes are synthetic,
built in temporary folders; they cover the layouts found in the real workbook."""
import contextlib
import io
import json
import os
import shutil
import sys
import tempfile
import unittest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import import_archive as ia  # noqa: E402

try:
    import openpyxl
    from openpyxl.styles import PatternFill
except ImportError:  # the Yahoo tests still run without it
    openpyxl = None

YELLOW, CYAN, RED = "FFFFFF00", "FF00FFFF", "FFFF0000"


def vault_note(pid, name, team="AAA", positions=("PG",)):
    return "---\nplayer_id: %s\nname: %s\nteam: %s\npositions: %r\nfantasy_teams: []\n---\n\n# %s\n" % (
        pid, name, team, list(positions), name)


def fill(cell, rgb):
    cell.fill = PatternFill(fill_type="solid", fgColor=rgb)


@unittest.skipIf(openpyxl is None, "openpyxl not installed")
class ImporterTest(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, self.dir)
        self.vault = os.path.join(self.dir, "vault")
        os.makedirs(os.path.join(self.vault, "Silmaril", "players"))
        os.makedirs(os.path.join(self.dir, "workbooks"))
        for pid, name in (("100", "Nikola Jokić"), ("200", "Jalen Williams"), ("300", "Tyrese Maxey"),
                          ("400", "Sam Smith"), ("401", "Sam  Smith"), ("500", "Bobby Portis")):
            slug = ia.normalize_name(name) + pid
            with open(os.path.join(self.vault, "Silmaril", "players", slug + ".md"), "w", encoding="utf-8") as f:
                f.write(vault_note(pid, name))
        with open(os.path.join(self.vault, "Silmaril", "Rules.md"), "w", encoding="utf-8") as f:
            f.write("[[index]]\n\n1. First rule\n\t- 1.a Nested **bold**\n        - 1.ab deeper\n"
                    "    - 1.b ~~old.*~~ new\n2. Second\n    continuation line\n\n[[Old Rules]]\n")
        with open(os.path.join(self.vault, "Silmaril", "Old Rules.md"), "w", encoding="utf-8") as f:
            f.write("1. ~~Struck~~ text _**both**_\n")
        self.aliases = os.path.join(self.dir, "aliases.json")
        self.write_json(self.aliases, {"aliases": [
            {"source_name": "Jalen Wiliams", "player_id": "200", "reason": "typo"}]})
        self.output = os.path.join(self.dir, "site", "archive.json")
        self.manifest_path = os.path.join(self.dir, "manifest.json")

    # ------------------------------------------------------------ helpers

    def write_json(self, path, data):
        with open(path, "w", encoding="utf-8") as f:
            json.dump(data, f)

    def workbook(self, draft_rows, keeper_rows=None, extra_sheet=False, name="league.xlsx"):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Draft"
        ws.append(["Pick", "Player", "Salary", "Team"])
        ws["F2"], ws["F3"] = "Consecutive Year Keeper Fee", "First time keeper fee"
        fill(ws["F2"], YELLOW)
        fill(ws["F3"], CYAN)
        for i, row in enumerate(draft_rows, start=2):
            values, link, colour = row[:4], (row[4] if len(row) > 4 else None), (row[5] if len(row) > 5 else None)
            for col, value in enumerate(values, start=1):
                ws.cell(row=i, column=col, value=value)
            if link:
                ws.cell(row=i, column=2).hyperlink = "https://sports.yahoo.com/nba/players/%s" % link
            if colour:
                fill(ws.cell(row=i, column=3), colour)
        if keeper_rows is not None:
            ks = wb.create_sheet("Keepers")
            ks.append(["Name", "Price to Keep"])
            for row in keeper_rows:
                ks.append(list(row))
        if extra_sheet:
            wb.create_sheet("Sheet3")["A1"] = "unmapped"
        wb.save(os.path.join(self.dir, "workbooks", name))

    def manifest(self, keeper=True, rosters=(), extra=None):
        sheets = [{"sheet": "Draft", "kind": "draft", "season": "2025-26", "header_row": 1,
                   "columns": {"pick": "Pick", "player": "Player", "salary": "Salary", "team": "Team"},
                   "player_name_format": "name_nba_team_positions",
                   "fee_type": {"season": "2026-27", "column": "salary", "legend_cells": ["F2", "F3"],
                                "extra_fills": {RED: "Can't Be Kept (CBK)"}},
                   "ignore_cells": ["F2", "F3"], "text_values": {"CBK": "Can't Be Kept"}}]
        if keeper:
            sheets.append({"sheet": "Keepers", "kind": "keeper", "season": "2025-26", "header_row": 1,
                           "columns": {"player": "Name", "cost": "Price to Keep"},
                           "team_from_draft": "2025-26", "text_values": {"CBK": "Can't Be Kept"}})
        if extra:
            sheets.extend(extra)
        self.write_json(self.manifest_path, {
            "vault": self.vault, "vault_players_dir": "Silmaril/players",
            "workbooks_dir": os.path.join(self.dir, "workbooks"),
            "workbooks": [{"file": "league.xlsx", "sheets": sheets}],
            "rosters": list(rosters),
            "rules": [{"id": "current", "file": "Silmaril/Rules.md", "title": "Rules", "status": "current"},
                      {"id": "old", "file": "Silmaril/Old Rules.md", "title": "Old", "status": "historical"}],
            "rules_links": {"index": {"unavailable": "vault home"}},
        })

    def run_import(self, *extra):
        out, err = io.StringIO(), io.StringIO()
        with contextlib.redirect_stdout(out), contextlib.redirect_stderr(err):
            code = ia.main(["--manifest", self.manifest_path, "--aliases", self.aliases,
                            "--output", self.output, "--report", os.path.join(self.dir, "REPORT.md"),
                            "--records", os.path.join(self.dir, "build", "records.json"),
                            "--now", "2026-10-09T00:00:00Z"] + list(extra))
        return code, out.getvalue() + err.getvalue()

    def report_text(self):
        with open(os.path.join(self.dir, "REPORT.md"), encoding="utf-8") as f:
            return f.read()

    def load(self):
        with open(self.output, encoding="utf-8") as f:
            return json.load(f)

    STANDARD = [
        (1, "Nikola Jokić (DEN - C)", 69, "Team A", "100", YELLOW),
        (2, "Tyrese Maxey (PHI - PG)", 0, "Team B", "300", CYAN),
        (3, "Jalen Williams (OKC - SF,PF)", "CBK", "Team B", "200", RED),
        (None, None, None, None),                             # blank row
        (4, "Bobby Portis Jr. (MIA - PF,C)", None, "Team A", "500", CYAN),
    ]

    # ------------------------------------------------------------ tests

    def test_imports_layout_values_and_matching(self):
        self.workbook(self.STANDARD, [("Nikola Jokic", 69), ("Jalen Wiliams", "CBK")])
        self.manifest()
        code, output = self.run_import()
        self.assertEqual(code, 0, output)
        data = self.load()
        drafts = {d["pick"]: d for d in data["drafts"]}
        self.assertEqual(len(drafts), 4)
        self.assertEqual(drafts[2]["salary"], 0)               # zero kept
        self.assertEqual(drafts[2]["salary_display"], "$0")
        self.assertIsNone(drafts[4]["salary"])                 # missing is not zero
        self.assertIsNone(drafts[4]["salary_display"])
        self.assertEqual(drafts[3]["salary_display"], "CBK")
        self.assertEqual(drafts[3]["salary_note"], "Can't Be Kept")
        self.assertEqual(drafts[1]["fee_type"], "Consecutive Year Keeper Fee")
        self.assertEqual(drafts[3]["fee_type"], "Can't Be Kept (CBK)")
        self.assertEqual((drafts[1]["nba_team"], drafts[3]["positions"]), ("DEN", ["SF", "PF"]))
        self.assertEqual(drafts[4]["player_id"], "500")        # link wins over the "Jr." name
        self.assertTrue(drafts[1]["kept"])
        self.assertFalse(drafts[2]["kept"])
        keepers = {k["player_id"]: k for k in data["keepers"]}
        self.assertEqual(keepers["100"]["team"], "Team A")     # team from the draft record
        self.assertEqual(keepers["100"]["draft_pick"], 1)
        self.assertEqual(keepers["200"]["cost_display"], "CBK")
        players = {p["id"]: p for p in data["players"]}
        self.assertIn("Bobby Portis Jr.", players["500"]["aliases"])
        self.assertIn("Jalen Wiliams", players["200"]["aliases"])
        self.assertEqual(data["seasons"]["draft"], ["2025-26"])
        self.assertEqual(data["imported_at"], "2026-10-09T00:00:00Z")
        report = self.report_text()
        self.assertIn("Skipped (blank row): 1: rows 5", report)
        self.assertIn("by explicit alias", report)
        self.assertIn("differs from vault name", report)
        # Local traceability stays out of the published file.
        self.assertNotIn("_source", json.dumps(data))
        self.assertNotIn(self.dir, json.dumps(data))
        with open(os.path.join(self.dir, "build", "records.json"), encoding="utf-8") as f:
            self.assertEqual(json.load(f)["d-2025-26-001"], [{"file": "league.xlsx", "sheet": "Draft", "row": 2}])

    def test_unmatched_and_ambiguous_names_get_stable_records(self):
        self.workbook([(1, "Unknown Guy (AAA - C)", 10, "Team A"), (2, "Sam Smith (AAA - C)", 10, "Team A")],
                      keeper_rows=[])
        self.manifest()
        code, output = self.run_import()
        self.assertEqual(code, 0, output)
        ids = [d["player_id"] for d in self.load()["drafts"]]
        self.assertEqual(ids, ["u-unknownguy", "u-samsmith"])   # never merged with 400 or 401
        report = self.report_text()
        self.assertIn("is ambiguous", report)
        self.assertIn("no vault match", report)

    def test_rerun_is_deterministic_and_keeps_import_date(self):
        self.workbook(self.STANDARD, [("Nikola Jokic", 69)])
        self.manifest()
        self.assertEqual(self.run_import()[0], 0)
        with open(self.output, "rb") as f:
            first = f.read()
        code, output = self.run_import("--now", "2030-01-01T00:00:00Z")
        self.assertEqual(code, 0)
        self.assertIn("Unchanged", output)
        with open(self.output, "rb") as f:
            self.assertEqual(f.read(), first)

    def test_failed_import_leaves_published_file_intact(self):
        self.workbook(self.STANDARD, [("Nikola Jokic", 69)])
        self.manifest()
        self.assertEqual(self.run_import()[0], 0)
        with open(self.output, "rb") as f:
            good = f.read()
        self.workbook([(1, "Nikola Jokić (DEN - C)", "lots", "Team A", "100")], [])   # invalid salary
        code, output = self.run_import()
        self.assertEqual(code, 1)
        self.assertIn("invalid amount", output)
        with open(self.output, "rb") as f:
            self.assertEqual(f.read(), good)
        self.assertIn("FAILED", self.report_text())

    def test_missing_required_team_is_rejected(self):
        self.workbook([(1, "Nikola Jokić (DEN - C)", 5, None, "100")], [])
        self.manifest()
        code, output = self.run_import()
        self.assertEqual(code, 1)
        self.assertIn("missing required Team", output)
        self.assertFalse(os.path.exists(self.output))

    def test_duplicate_pick_blocks_and_duplicate_player_warns(self):
        self.workbook([(1, "Nikola Jokić (DEN - C)", 5, "Team A", "100"),
                       (1, "Tyrese Maxey (PHI - PG)", 5, "Team B", "300")], [])
        self.manifest()
        code, output = self.run_import()
        self.assertEqual(code, 1)
        self.assertIn("draft pick 1 appears twice", output)

        self.workbook([(1, "Nikola Jokić (DEN - C)", 5, "Team A", "100"),
                       (2, "Nikola Jokić (DEN - C)", 6, "Team B", "100")], [])
        code, output = self.run_import()
        self.assertEqual(code, 0, output)
        self.assertIn("Possible duplicate", self.report_text())

    def test_keeper_without_draft_record_is_rejected(self):
        self.workbook(self.STANDARD, [("Sam Smith", 5)])
        self.manifest()
        code, output = self.run_import()
        self.assertEqual(code, 1)
        self.assertIn("draft record", output)

    def test_unlisted_sheet_stops_import(self):
        self.workbook(self.STANDARD, [], extra_sheet=True)
        self.manifest()
        code, output = self.run_import()
        self.assertEqual(code, 1)
        self.assertIn("'Sheet3' are not in the manifest", output)
        self.manifest(extra=[{"sheet": "Sheet3", "kind": "ignore", "reason": "test"}])
        self.assertEqual(self.run_import()[0], 0)

    def test_roster_team_blocks_and_snapshot_preservation(self):
        self.workbook(self.STANDARD, [])
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Rosters"
        for coord, value in {"A1": "Team A", "A2": "Player", "A3": "Nikola Jokić", "A4": "Neemias",
                             "B1": "Team B", "B2": "Player", "B3": "Tyrese Maxey"}.items():
            ws[coord] = value
        ws["A4"].hyperlink = "https://sports.yahoo.com/nba/players/300"
        wb.save(os.path.join(self.dir, "workbooks", "rosters.xlsx"))
        roster = {"file": "rosters.xlsx", "sheet": "Rosters", "kind": "roster", "layout": "team_blocks",
                  "season": "2025-26", "snapshot_date": "2026-04-12"}
        self.manifest(rosters=[roster])
        code, output = self.run_import()
        self.assertEqual(code, 0, output)
        snap = self.load()["rosters"][0]
        self.assertEqual(snap["id"], "r-2025-26-2026-04-12")
        self.assertEqual([(e["team"], e["player_id"]) for e in snap["entries"]],
                         [("Team A", "300"), ("Team A", "100"), ("Team B", "300")])
        # Dropping a published snapshot from the manifest is refused.
        self.manifest(rosters=[])
        code, output = self.run_import()
        self.assertEqual(code, 1)
        self.assertIn("already published", output)
        self.assertEqual(len(self.load()["rosters"]), 1)

    def test_roster_requires_snapshot_date(self):
        self.workbook(self.STANDARD, [])
        self.manifest(rosters=[{"file": "league.xlsx", "sheet": "Draft", "kind": "roster",
                                "layout": "table", "season": "2025-26",
                                "columns": {"team": "Team", "player": "Player"}}])
        code, output = self.run_import()
        self.assertEqual(code, 1)
        self.assertIn("snapshot_date", output)

    def test_rules_are_generated_from_vault(self):
        self.workbook(self.STANDARD, [])
        self.manifest()
        self.assertEqual(self.run_import()[0], 0)
        rules = {d["id"]: d for d in self.load()["rules"]}
        blocks = rules["current"]["blocks"]
        self.assertEqual(blocks[0]["c"][0], {"t": "missing", "v": "index", "note": "vault home"})
        items = blocks[1]["items"]
        self.assertEqual([i["n"] for i in items], [1, 2])
        nested = items[0]["blocks"][0]["items"]
        self.assertEqual(len(nested), 2)                       # 1.a (tab) and 1.b (4 spaces)
        self.assertEqual(nested[0]["blocks"][0]["items"][0]["c"], [{"t": "text", "v": "1.ab deeper"}])
        self.assertEqual(nested[1]["c"][1], {"t": "del", "c": [{"t": "text", "v": "old.*"}]})
        self.assertEqual(items[1]["blocks"][0]["t"], "p")      # indented continuation
        self.assertEqual(blocks[2]["c"][0]["doc"], "old")
        self.assertEqual(rules["old"]["blocks"][0]["items"][0]["c"][2],
                         {"t": "em", "c": [{"t": "strong", "c": [{"t": "text", "v": "both"}]}]})

    # -------------------------------------------- draft prices, roster prices, accuracy

    def priced_setup(self, price_rows, keeper_rows=(), roster=None):
        """Draft sheet with next-season prices and fee colours, plus a separate draft-price sheet."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Draft"
        ws.append(["Pick", "Player", "Salary", "Team"])
        ws["F2"], ws["F3"], ws["F4"] = "Consecutive Year Keeper Fee", "First time keeper fee", "Drafted for Sub $10 fee"
        fill(ws["F2"], YELLOW)
        fill(ws["F3"], CYAN)
        fill(ws["F4"], "FFB6D7A8")
        draft = [(1, "Nikola Jokić (DEN - C)", 69, "Team A", "100", YELLOW),
                 (2, "Tyrese Maxey (PHI - PG)", 59, "Team B", "300", CYAN),
                 (3, "Jalen Williams (OKC - SF)", 15, "Team B", "200", "FFB6D7A8")]
        for i, (pick, name, price, team, pid, colour) in enumerate(draft, start=2):
            for col, value in enumerate((pick, name, price, team), start=1):
                ws.cell(row=i, column=col, value=value)
            ws.cell(row=i, column=2).hyperlink = "https://sports.yahoo.com/nba/players/%s" % pid
            fill(ws.cell(row=i, column=3), colour)
        ks = wb.create_sheet("Keepers")
        ks.append(["Name", "Price to Keep"])
        for row in keeper_rows:
            ks.append(list(row))
        rs = wb.create_sheet("Rosters")
        for coord, value in (roster or {}).items():
            rs[coord] = value
        wb.save(os.path.join(self.dir, "workbooks", "league.xlsx"))
        pw = openpyxl.Workbook()
        ps = pw.active
        ps.title = "Prices"
        ps.append(["Pick", "Player", "Salary", "Team"])
        for i, (pick, name, price, pid) in enumerate(price_rows, start=2):
            for col, value in enumerate((pick, name, price), start=1):
                ps.cell(row=i, column=col, value=value)
            ps.cell(row=i, column=2).hyperlink = "https://sports.yahoo.com/nba/players/%s" % pid
        pw.save(os.path.join(self.dir, "workbooks", "prices.xlsx"))
        sheets = [
            {"sheet": "Draft", "kind": "draft", "season": "2025-26",
             "columns": {"pick": "Pick", "player": "Player", "next_price": "Salary", "team": "Team"},
             "player_name_format": "name_nba_team_positions",
             "fee_type": {"season": "2026-27", "column": "next_price", "legend_cells": ["F2", "F3", "F4"],
                          "display": {"First time keeper fee": {"key": "first", "label": "1st-year keeper", "short": "1st"},
                                      "Consecutive Year Keeper Fee": {"key": "repeat", "label": "Repeat keeper"},
                                      "Drafted for Sub $10 fee": {"key": "under10", "label": "Under $10"}}},
             "ignore_cells": ["F2", "F3", "F4"]},
            {"sheet": "Keepers", "kind": "keeper", "season": "2025-26",
             "columns": {"player": "Name", "next_cost": "Price to Keep"},
             "team_from_draft": "2025-26", "from_draft": ["cost"]},
            ({"sheet": "Rosters", "kind": "roster", "layout": "team_blocks", "season": "2026-27",
              "label": "Pre-draft", "snapshot_date": "2026-10-09", "keeper_prices_from_draft": "2025-26",
              "undrafted_price": {"amount": 15, "basis": "rule 1.c", "fee_key": "undrafted", "fee_short": "Undrafted"}}
             if roster else {"sheet": "Rosters", "kind": "ignore", "reason": "none"}),
        ]
        self.write_json(self.manifest_path, {
            "vault": self.vault, "workbooks_dir": os.path.join(self.dir, "workbooks"),
            "workbooks": [{"file": "league.xlsx", "sheets": sheets},
                          {"file": "prices.xlsx", "sheets": [
                              {"sheet": "Prices", "kind": "draft_prices", "season": "2025-26",
                               "columns": {"pick": "Pick", "player": "Player", "salary": "Salary"},
                               "player_name_format": "name_nba_team_positions", "ignore_cells": ["D1"]}]}],
            "rules": [],
            "accuracy_check": {"season": "2025-26", "rules": {
                "first": {"salary_at_least": 10, "min_price": 10, "pct": 10, "min_fee": 5},
                "under10": {"salary_below": 10, "min_price": 10, "pct": 10, "min_fee": 5},
                "repeat": {"min_price": 10, "pct": 10, "min_fee": 7}}},
        })

    GOOD_PRICES = [(1, "Nikola Jokić (DEN - C)", 62, "100"), (2, "Tyrese Maxey (PHI - PG)", 54, "300"),
                   (3, "Jalen Williams (OKC - SF)", 1, "200")]

    def test_draft_prices_merge_and_keeper_cost_from_draft(self):
        self.priced_setup(self.GOOD_PRICES, keeper_rows=[("Nikola Jokic", 69)])
        code, output = self.run_import()
        self.assertEqual(code, 0, output)
        data = self.load()
        d = {r["pick"]: r for r in data["drafts"]}
        self.assertEqual((d[1]["salary"], d[1]["next_price"]), (62, 69))
        self.assertEqual((d[3]["salary_display"], d[3]["next_price_display"]), ("$1", "$15"))
        self.assertEqual((d[2]["fee_key"], d[2]["fee_type"], d[2]["fee_short"]), ("first", "1st-year keeper", "1st"))
        self.assertEqual((d[1]["fee_short"]), "Repeat keeper")       # short falls back to the label
        k = data["keepers"][0]
        self.assertEqual((k["cost"], k["next_cost"], k["team"], k["draft_pick"]), (62, 69, "Team A", 1))
        report = self.report_text()
        self.assertIn("Checked: 3. Correct: 3. Wrong: 0.", report)
        self.assertIn("Pick 2 Tyrese Maxey: draft price $54, recorded $59; 10% is $5.40, rounded to $5", report)
        with open(os.path.join(self.dir, "build", "records.json"), encoding="utf-8") as f:
            self.assertEqual(json.load(f)["d-2025-26-001"][1]["file"], "prices.xlsx")

    def test_accuracy_reports_wrong_prices_without_blocking(self):
        self.priced_setup([(1, "Nikola Jokić (DEN - C)", 60, "100"), (2, "Tyrese Maxey (PHI - PG)", 54, "300"),
                           (3, "Jalen Williams (OKC - SF)", 12, "200")])
        code, output = self.run_import()
        self.assertEqual(code, 0, output)
        report = self.report_text()
        self.assertIn("Wrong: 2.", report)
        self.assertIn("| 1 | Nikola Jokić | $60 | $69 | Consecutive Year Keeper Fee | expected $67 |", report)
        self.assertIn("fee colour does not fit a $12 draft price", report)
        self.assertNotIn("accuracy", json.dumps(self.load()).lower())   # report only

    def test_half_dollar_fee_rounds_up(self):
        # $55 first-time keeper: 10% is $5.50, which rounds up to $6, so $61 is right and $60 is wrong.
        self.priced_setup([(1, "Nikola Jokić (DEN - C)", 62, "100"), (2, "Tyrese Maxey (PHI - PG)", 55, "300"),
                           (3, "Jalen Williams (OKC - SF)", 1, "200")])
        self.assertEqual(self.run_import()[0], 0)
        self.assertIn("| 2 | Tyrese Maxey | $55 | $59 | First time keeper fee | expected $61 |", self.report_text())

    def test_draft_price_for_wrong_player_is_rejected(self):
        self.priced_setup([(1, "Tyrese Maxey (PHI - PG)", 62, "300")])
        code, output = self.run_import()
        self.assertEqual(code, 1)
        self.assertIn("pick 1 is Tyrese Maxey here but Nikola Jokić in the draft sheet", output)
        self.assertFalse(os.path.exists(self.output))

    def test_roster_prices_drafted_moved_and_undrafted(self):
        self.priced_setup(self.GOOD_PRICES, roster={
            "A1": "Team A", "A2": "Player", "A3": "Tyrese Maxey", "A4": "Bobby Portis", "A5": "Nobody Known",
            "B1": "Team B", "B2": "Player", "B3": "Jalen Williams"})
        code, output = self.run_import()
        self.assertEqual(code, 0, output)
        snap = self.load()["rosters"][0]
        self.assertEqual((snap["label"], snap["price_from_draft"], snap["undrafted_rule"]),
                         ("Pre-draft", "2025-26", "rule 1.c"))
        e = {x["source_name"]: x for x in snap["entries"]}
        self.assertEqual((e["Tyrese Maxey"]["keeper_price"], e["Tyrese Maxey"]["drafted_by"]), (59, "Team B"))
        self.assertEqual((e["Bobby Portis"]["keeper_price_display"], e["Bobby Portis"]["price_basis"]), ("$15", "undrafted"))
        self.assertEqual(e["Nobody Known"]["price_basis"], "undrafted")
        self.assertEqual(e["Jalen Williams"]["keeper_price"], 15)
        self.assertEqual((e["Tyrese Maxey"]["fee_key"], e["Bobby Portis"]["fee_key"]), ("first", "undrafted"))

    def test_unmapped_fee_label_stops_import(self):
        self.priced_setup(self.GOOD_PRICES)
        with open(self.manifest_path, encoding="utf-8") as f:
            manifest = json.load(f)
        del manifest["workbooks"][0]["sheets"][0]["fee_type"]["display"]["Drafted for Sub $10 fee"]
        self.write_json(self.manifest_path, manifest)
        code, output = self.run_import()
        self.assertEqual(code, 1)
        self.assertIn("have no entry in fee_type.display", output)

    def test_rules_link_to_site_section_and_suggestion_form(self):
        self.workbook(self.STANDARD, [])
        self.manifest()
        with open(self.manifest_path, encoding="utf-8") as f:
            manifest = json.load(f)
        manifest["rules_links"]["old rules"] = {"section": "suggestions"}
        manifest["rules"] = manifest["rules"][:1]
        manifest["suggestions"] = {"title": "Suggest", "season": "2026-27", "form_url": "https://forms.gle/abc123"}
        self.write_json(self.manifest_path, manifest)
        code, output = self.run_import()
        self.assertEqual(code, 0, output)
        data = self.load()
        self.assertEqual(data["rules"][0]["blocks"][2]["c"][0],
                         {"t": "link", "section": "suggestions", "c": [{"t": "text", "v": "Old Rules"}]})
        self.assertEqual(data["suggestions"]["form_url"], "https://forms.gle/abc123")
        manifest["suggestions"]["form_url"] = "javascript:alert(1)"
        self.write_json(self.manifest_path, manifest)
        code, output = self.run_import()
        self.assertEqual(code, 1)
        self.assertIn("must be a Google Forms link", output)


class FormulaHandlingTest(unittest.TestCase):
    """The importer uses Excel's saved result and flags formulas without one."""

    class FakeCell:
        def __init__(self, value):
            self.value = value

    class FakeSheet:
        def __init__(self, values):
            self.values = values

        def cell(self, row, column):
            return FormulaHandlingTest.FakeCell(self.values.get((row, column)))

    def test_saved_result_used_and_missing_result_flagged(self):
        sheet = ia.Sheet(self.FakeSheet({(1, 1): "=B1*2", (2, 1): "=B2*2", (3, 1): 0}),
                         self.FakeSheet({(1, 1): 44, (2, 1): None, (3, 1): 0}))
        self.assertEqual(sheet.value(1, 1), (44, None))
        value, problem = sheet.value(2, 1)
        self.assertIsNone(value)
        self.assertIn("no saved result", problem)
        self.assertEqual(sheet.value(3, 1), (0, None))         # a real zero is not a problem


class HelperTest(unittest.TestCase):
    def test_money(self):
        self.assertEqual(ia.parse_money(22.0, {})[:2], (22, "$22"))
        self.assertEqual(ia.parse_money("$15", {})[:2], (15, "$15"))
        self.assertEqual(ia.parse_money(0, {})[:2], (0, "$0"))
        self.assertEqual(ia.parse_money(None, {}), (None, None, None, None))
        self.assertEqual(ia.parse_money("CBK", {"CBK": "Can't Be Kept"})[:3], (None, "CBK", "Can't Be Kept"))
        self.assertIn("invalid", ia.parse_money("ten", {})[3])
        self.assertIn("negative", ia.parse_money(-1, {})[3])

    def test_season(self):
        for label in ("2025-26", "'25-'26", "2025–2026", "25-26"):
            self.assertEqual(ia.normalize_season(label), "2025-26")
        with self.assertRaises(ValueError):
            ia.normalize_season("2025-27")

    def test_name_normalisation(self):
        self.assertEqual(ia.normalize_name("Alperen Şengün"), ia.normalize_name("Alperen Sengun"))
        self.assertEqual(ia.normalize_name("Kel'el Ware"), "kelelware")
        self.assertNotEqual(ia.normalize_name("Bobby Portis Jr."), ia.normalize_name("Bobby Portis"))


if __name__ == "__main__":
    unittest.main()
